"""Plano de migração em planilha (PLANO_ESCOPO_DE_MIGRACAO, E7).

O arquivo é o mesmo plano que a tela mostra, numa aba por pergunta: o que vai
junto, o que precisa ir antes, quem fica para trás, quanto custa e em que ordem.
Só formata; nenhum número é calculado aqui.
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from app.schemas.migration_scope import MigrationPlanResponse

_TITULO = Font(bold=True, size=14)
_CABECALHO = Font(bold=True, color="FFFFFF")
_FUNDO = PatternFill("solid", fgColor="644AFF")
_DINHEIRO = "#,##0.00"
_PORCENTAGEM = "0.0%"
_NATUREZA = {
    "dados": "acessa dados",
    "servico": "chama serviço",
    "encaminhamento": "encaminha tráfego",
    "declarada": "declarada",
    "administrativo": "administrativo",
}
_METODO = {
    "pool": "pool da família e região",
    "pool_da_familia": "média da família",
    "media_dos_pools": "média dos pools",
    "categoria": "razão da categoria",
}


def _ordem(valor: int | None) -> str:
    return str(valor) if valor is not None else "livre"


def _tabela(
    ws: Any,
    linha: int,
    cabecalho: Sequence[str],
    linhas: Iterable[Sequence[Any]],
    dinheiro: Iterable[int] = (),
) -> int:
    for col, texto in enumerate(cabecalho, start=1):
        c = ws.cell(row=linha, column=col, value=texto)
        c.font, c.fill = _CABECALHO, _FUNDO
    colunas_dinheiro = set(dinheiro)
    for valores in linhas:
        linha += 1
        for col, valor in enumerate(valores, start=1):
            c = ws.cell(row=linha, column=col, value=valor)
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(valor, str))
            if col in colunas_dinheiro and isinstance(valor, int | float):
                c.number_format = _DINHEIRO
    return linha + 2


def _larguras(ws: Any, larguras: Sequence[int]) -> None:
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura


def montar_xlsx(plano: MigrationPlanResponse, *, job_id: str, gerado_em: datetime) -> bytes:
    wb = Workbook()
    moeda = plano.moeda
    escolha = plano.escolha

    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "Plano de migração"
    ws["A1"].font = _TITULO
    info = [
        ("Análise", job_id),
        ("Gerado em", gerado_em.strftime("%d/%m/%Y %H:%M UTC")),
        ("Nuvem atual", (plano.provedor_origem or "").upper()),
        ("Recomendada pela análise", (plano.provedor_recomendado or "").upper()),
        ("Aplicações escolhidas", ", ".join(escolha.aplicacoes) or "-"),
        ("Serviços escolhidos", ", ".join(escolha.servicos) or "-"),
        ("Ambiente", escolha.ambiente or "todos"),
        ("Grupos incluídos à mão", ", ".join(escolha.incluir_grupos) or "-"),
        ("Aplicações no escopo", ", ".join(plano.aplicacoes)),
        ("Recursos no inventário", plano.recursos),
        (f"Custo hoje ({moeda}/mês)", plano.custo),
    ]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        c = ws.cell(row=i, column=2, value=v)
        if isinstance(v, float):
            c.number_format = _DINHEIRO
    linha = _tabela(
        ws,
        len(info) + 4,
        ["Destino", f"Estimado ({moeda}/mês)", "Origem com cotação", "Variação", "Sem cotação"],
        (
            [
                d.provedor.upper(),
                d.custo_estimado,
                d.custo_origem_cotado,
                d.variacao,
                d.sem_cotacao,
            ]
            for d in plano.destino
        ),
        dinheiro=(2, 3, 5),
    )
    for row in ws.iter_rows(min_row=len(info) + 5, max_row=linha, min_col=4, max_col=4):
        for c in row:
            if isinstance(c.value, float):
                c.number_format = _PORCENTAGEM
    ws.cell(row=linha, column=1, value="Premissas").font = Font(bold=True)
    for i, p in enumerate([plano.limite_das_dependencias, *(x.base for x in plano.premissas)]):
        ws.cell(row=linha + 1 + i, column=1, value=p).alignment = Alignment(wrap_text=False)
    _larguras(ws, [34, 60, 20, 12, 16])

    ws = wb.create_sheet("Vão junto")
    _tabela(
        ws,
        1,
        ["Unidade", "Aplicação", "Por quê", f"Custo ({moeda}/mês)", "Incluída por você"],
        (
            [v.nome, v.aplicacao, v.motivo, v.custo, "sim" if v.incluida_por_voce else "não"]
            for v in plano.vao_junto
        ),
        dinheiro=(4,),
    )
    _larguras(ws, [36, 18, 60, 16, 16])
    if plano.servicos_que_vem_junto:
        inicio = len(plano.vao_junto) + 4
        ws.cell(row=inicio - 1, column=1, value="Serviços que vêm junto").font = Font(bold=True)
        _tabela(
            ws,
            inicio,
            ["Serviço", "Por quê", f"Custo ({moeda}/mês)"],
            ([s.servico, s.motivo, s.custo] for s in plano.servicos_que_vem_junto),
            dinheiro=(3,),
        )

    for titulo, grupos, explicacao in (
        (
            "Antes",
            plano.antes,
            "O escopo chama estes grupos: migre antes ou mantenha o acesso entre nuvens.",
        ),
        (
            "Para trás",
            plano.para_tras,
            "Estes grupos chamam o escopo: vão acessar entre nuvens (latência e tráfego de saída).",
        ),
    ):
        ws = wb.create_sheet(titulo)
        ws["A1"] = explicacao
        _tabela(
            ws,
            3,
            ["Grupo", "Ordem", "Ligação", "Portas", f"Custo ({moeda}/mês)", "Nos dois sentidos"],
            (
                [
                    g.nome,
                    _ordem(g.ordem),
                    ", ".join(_NATUREZA.get(n, n) for n in g.naturezas),
                    ", ".join(g.portas),
                    g.custo,
                    "sim" if g.nos_dois_sentidos else "não",
                ]
                for g in grupos
            ),
            dinheiro=(5,),
        )
        _larguras(ws, [48, 8, 26, 22, 16, 16])

    ws = wb.create_sheet("Serviços")
    provedores = [d.provedor for d in plano.destino]
    _tabela(
        ws,
        1,
        ["Serviço", f"Hoje ({moeda}/mês)", *(p.upper() for p in provedores), "Método"],
        (
            [
                s.nome,
                s.custo,
                *(s.destino.get(p) for p in provedores),
                ", ".join(_METODO.get(m, m) for m in s.metodos) or "sem cotação",
            ]
            for s in plano.servicos
        ),
        dinheiro=range(2, 3 + len(provedores)),
    )
    _larguras(ws, [44, 16, *([14] * len(provedores)), 36])

    ws = wb.create_sheet("Ligações")
    ws["A1"] = "Como os serviços do escopo se chamam (das regras de rede)."
    _tabela(
        ws,
        3,
        ["Serviço", "Ligação", "Serviço chamado", "Portas"],
        (
            [x.de, _NATUREZA.get(x.natureza, x.natureza), x.para, ", ".join(x.portas)]
            for x in plano.ligacoes
        ),
    )
    _larguras(ws, [44, 20, 44, 22])

    ws = wb.create_sheet("Ordem")
    ws["A1"] = "Ordem de migração: a 1 vai primeiro; livre não depende de ninguém."
    _tabela(
        ws, 3, ["Ordem", "Grupos"], ([_ordem(o.ordem), " · ".join(o.grupos)] for o in plano.ordens)
    )
    _larguras(ws, [10, 120])

    saida = BytesIO()
    wb.save(saida)
    return saida.getvalue()
