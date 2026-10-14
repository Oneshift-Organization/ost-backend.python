"""Exportação do plano de migração em planilha (PLANO_ESCOPO_DE_MIGRACAO, E7)."""

from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.infrastructure.exports.plano_xlsx import montar_xlsx
from app.schemas.migration_scope import (
    EscolhaOut,
    GrupoVizinhoOut,
    LigacaoOut,
    MigrationPlanResponse,
    OrdemOut,
    ServicoNoEscopoOut,
    ServicoQueVemJuntoOut,
    TotalNoDestinoOut,
    UnidadeQueVaiJuntoOut,
)


def _plano() -> MigrationPlanResponse:
    return MigrationPlanResponse(
        moeda="USD",
        provedor_origem="aws",
        provedor_recomendado="oci",
        escolha=EscolhaOut(aplicacoes=[], servicos=["EC2"], ambiente="prod", incluir_grupos=[]),
        aplicacoes=["auth", "checkout"],
        unidades=2,
        recursos=7,
        custo=700.0,
        servicos=[
            ServicoNoEscopoOut(
                nome="RDS", custo=550.0, destino={"oci": 275.0}, metodos=["categoria"]
            ),
            ServicoNoEscopoOut(nome="ELB", custo=10.0, destino={"oci": None}, metodos=[]),
        ],
        destino=[
            TotalNoDestinoOut(
                provedor="oci",
                custo_estimado=350.0,
                custo_origem_cotado=690.0,
                sem_cotacao=10.0,
                variacao=-0.4928,
            )
        ],
        vao_junto=[
            UnidadeQueVaiJuntoOut(
                unidade="u1",
                nome="auth",
                aplicacao="auth",
                custo=500.0,
                motivo="Compartilha banco com checkout (tcp/5432)",
                incluida_por_voce=False,
                grupo="g1",
            )
        ],
        servicos_que_vem_junto=[
            ServicoQueVemJuntoOut(servico="RDS", custo=550.0, motivo="EC2 acessa dados em RDS")
        ],
        antes=[
            GrupoVizinhoOut(
                grupo="g2",
                nome="catalogo",
                ordem=1,
                naturezas=["servico"],
                portas=["tcp/8080"],
                custo=100.0,
                nos_dois_sentidos=False,
            )
        ],
        para_tras=[],
        ligacoes=[LigacaoOut(de="EC2", para="RDS", natureza="dados", portas=["tcp/5432"])],
        ordens=[OrdemOut(ordem=2, grupos=["auth + checkout"]), OrdemOut(ordem=None, grupos=["x"])],
        fora_do_plano=[],
        premissas=[],
    )


def test_planilha_tem_uma_aba_por_pergunta_do_plano() -> None:
    conteudo = montar_xlsx(_plano(), job_id="j-1", gerado_em=datetime(2026, 10, 13, tzinfo=UTC))
    wb = load_workbook(BytesIO(conteudo))
    assert wb.sheetnames == [
        "Resumo",
        "Vão junto",
        "Antes",
        "Para trás",
        "Serviços",
        "Ligações",
        "Ordem",
    ]

    def valores(aba: str) -> list[list[object]]:
        return [
            [v for v in linha if v is not None]
            for linha in wb[aba].iter_rows(values_only=True)
            if any(v is not None for v in linha)
        ]

    resumo = valores("Resumo")
    assert ["Serviços escolhidos", "EC2"] in resumo
    assert ["OCI", 350.0, 690.0, -0.4928, 10.0] in resumo
    assert valores("Vão junto")[1][2] == "Compartilha banco com checkout (tcp/5432)"
    assert ["RDS", "EC2 acessa dados em RDS", 550.0] in valores("Vão junto")
    assert ["catalogo", "1", "chama serviço", "tcp/8080", 100.0, "não"] in valores("Antes")
    # Sem cotação é dito, não deixado em branco
    assert ["ELB", 10.0, "sem cotação"] in valores("Serviços")
    assert ["EC2", "acessa dados", "RDS", "tcp/5432"] in valores("Ligações")
    assert ["livre", "x"] in valores("Ordem")
