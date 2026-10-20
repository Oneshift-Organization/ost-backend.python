"""Plano de migração exportado: documento, PDF e planilha (PLANO_ESCOPO_DE_MIGRACAO, E7)."""

from __future__ import annotations

from datetime import UTC, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook  # type: ignore[import-untyped]

from app.domain.entities.topologia import TipoDeRecurso
from app.domain.services.correspondencia_inventario import Correspondencia
from app.domain.services.cotacao_do_escopo import LinhaDoEscopo, PoolCotado, cotar_escopo
from app.domain.services.topologia.escopo import Escolha, fechar_plano
from app.infrastructure.exports.plano_documento import (
    POS,
    PRE,
    PRINCIPAL,
    Documento,
    contagem,
    montar_documento,
)
from app.infrastructure.exports.plano_pdf import montar_pdf
from app.infrastructure.exports.plano_xlsx import montar_xlsx
from app.schemas.migration_scope import MigrationPlanResponse
from app.tests.unit.domain.topologia.test_escopo_de_migracao import _CUSTOS, _cenario
from app.use_cases.analysis.escopo_de_migracao import (
    ContextoDoEscopo,
    PlanoCotado,
    RecursoDoEscopo,
)

D = Decimal


def _documento(*, incluir_antes: bool, incluir_depois: bool) -> Documento:
    catalogo = _cenario()
    analise = catalogo.analise
    base = fechar_plano(catalogo, Escolha(aplicacoes=frozenset({"checkout"})))
    incluir = set()
    if incluir_antes:
        incluir.add(base.antes[0].grupo)
    if incluir_depois:
        incluir.add(base.para_tras[0].grupo)
    plano = fechar_plano(
        catalogo,
        Escolha(aplicacoes=frozenset({"checkout"}), incluir_grupos=frozenset(incluir)),
    )
    linhas = [
        LinhaDoEscopo(
            servico=servico,
            custo=D(custo),
            tipo=TipoDeRecurso.INSTANCIA if chave.startswith("vm") else TipoDeRecurso.BANCO,
            familia="M7G" if chave.startswith("vm") else None,
            regiao="us-east-1",
            unidade=analise.unidade_de[chave],
            recurso=chave,
        )
        for chave, (servico, custo) in _CUSTOS.items()
    ]
    contexto = ContextoDoEscopo(
        catalogo=catalogo,
        linhas=tuple(linhas),
        pools=(PoolCotado("M7G", "us-east-1", D("100"), {"oci": D("50")}),),
        taxas_de_categoria={"oci": {"banco": D("0.5")}},
        provedores=("oci",),
        provedor_origem="aws",
        provedor_recomendado="oci",
        moeda="USD",
        inventario_id="inv-1",
        correspondencia=Correspondencia(
            status="ok",
            fracao_encontrada=D("1"),
            custo_com_recurso=D("815"),
            custo_encontrado=D("815"),
            contas_da_fatura=("111",),
            contas_do_inventario=("111",),
            contas_sem_inventario=(),
            mensagem="ok",
        ),
        premissas=(
            ("chave_aplicacao", "app", "Chave de aplicação detectada nas tags."),
            ("x", None, "Chave de aplicação detectada nas tags."),
        ),
        recursos={
            chave: RecursoDoEscopo(
                chave=f"arn:aws:ec2:us-east-1:111:{chave}",
                nome=chave,
                tipo="instancia" if chave.startswith("vm") else "banco",
                tipo_nativo="AWS::EC2::Instance"
                if chave.startswith("vm")
                else "AWS::RDS::DBInstance",
                regiao="us-east-1",
                custo=D(_CUSTOS[chave][1]),
            )
            for chave in _CUSTOS
        },
        periodo="2026-03",
        gasto_total=D("815"),
    )
    linhas_do_plano = [x for x in linhas if x.unidade in set(plano.unidades)]
    cotado = PlanoCotado(
        plano=plano,
        cotacao=cotar_escopo(
            linhas_do_plano, contexto.pools, contexto.taxas_de_categoria, ("oci",)
        ),
        contexto=contexto,
    )
    return montar_documento(
        MigrationPlanResponse.de(cotado),
        cotado,
        job_id="job-1",
        gerado_em=datetime(2026, 10, 14, 16, 30, tzinfo=UTC),
    )


def test_documento_separa_pre_principal_e_pos_como_a_tela() -> None:
    doc = _documento(incluir_antes=True, incluir_depois=True)
    assert [b.chave for b in doc.blocos] == [PRE, PRINCIPAL, POS]
    pre, principal, pos = doc.blocos
    assert [(i.nome, i.papel) for i in pre.itens] == [("catalogo", "pré-requisito")]
    assert pre.itens[0].detalhe.startswith("Chamado por ")
    assert [i.papel for i in principal.itens] == ["escolhido", "vai junto"]
    # Passos numerados do começo ao fim do plano, um por bloco e etapa
    assert [(b.chave, p.numero, p.ordem) for b in doc.blocos for p in b.passos] == [
        (PRE, 1, pre.itens[0].ordem),
        (PRINCIPAL, 2, principal.itens[0].ordem),
        (POS, 3, pos.itens[0].ordem),
    ]
    assert principal.itens[1].detalhe.startswith("Compartilha banco com")
    assert [(i.nome, i.papel) for i in pos.itens] == [("monitor", "usa o plano")]
    assert pos.itens[0].detalhe.startswith("Chama ")
    assert (pre.custo, principal.custo, pos.custo) == (100.0, 700.0, 15.0)
    # Tudo incluído: nada fora do plano
    assert doc.pendencias == ()
    # Cada recurso no bloco do grupo dele, na ordem do plano
    assert doc.recursos[0].bloco == PRE
    assert {r.bloco for r in doc.recursos if r.nome.endswith("monitor")} == {POS}
    # Premissa repetida aparece uma vez
    assert len(doc.premissas) == len(set(doc.premissas))


def test_sem_inclusao_o_que_conversa_com_o_plano_vira_pendencia_e_a_ordem_explica() -> None:
    doc = _documento(incluir_antes=False, incluir_depois=False)
    assert [b.chave for b in doc.blocos] == [PRINCIPAL]
    assert [(p.nome, p.lado) for p in doc.pendencias] == [
        ("catalogo", "antes"),
        ("monitor", "depois"),
    ]
    assert all(p.ligado_a and p.acao for p in doc.pendencias)
    # O plano começa depois da etapa do catalogo: o aviso diz o que ficou de fora
    assert doc.aviso_da_ordem is not None
    assert "catalogo" in doc.aviso_da_ordem
    fora = [g.nome for e in doc.etapas for g in e.grupos if g.bloco is None]
    assert set(fora) == {"catalogo", "monitor"}


def test_pdf_sai_com_o_plano() -> None:
    conteudo = montar_pdf(_documento(incluir_antes=True, incluir_depois=False))
    assert conteudo.startswith(b"%PDF")
    assert len(conteudo) > 3000


def test_planilha_e_a_lista_de_trabalho_de_quem_executa() -> None:
    doc = _documento(incluir_antes=True, incluir_depois=False)
    wb = load_workbook(BytesIO(montar_xlsx(doc)))
    assert wb.sheetnames == [
        "Resumo",
        "Plano",
        "Recursos",
        "Fora do plano",
        "Ordem",
        "Serviços",
        "Ligações",
    ]

    def valores(aba: str) -> list[list[object]]:
        return [
            [v for v in linha if v is not None]
            for linha in wb[aba].iter_rows(values_only=True)
            if any(v is not None for v in linha)
        ]

    resumo = valores("Resumo")
    assert resumo[0] == ["Plano de migração: checkout"]
    # Nomes, nunca identificadores de grupo
    texto = str(resumo)
    assert "arn:" not in texto
    assert ["Pré-migração", "catalogo", 100.0] in resumo
    plano = valores("Plano")
    assert plano[2][:5] == [1, "Pré-migração", "1ª etapa", "catalogo", "pré-requisito"]
    recursos = valores("Recursos")
    assert recursos[1][-3:] == ["Status", "Responsável", "Observações"]
    assert all(linha[-1] == "A migrar" for linha in recursos[2:])
    assert wb["Recursos"].data_validations.dataValidation
    assert [linha[0] for linha in valores("Fora do plano")[2:]] == ["monitor"]
    servicos = valores("Serviços")
    assert ["EC2", 180, 90, "pool da família e região"] in servicos


def test_contagem_por_tipo() -> None:
    assert contagem("Instância", 1) == "1 instância"
    assert contagem("Instância", 14) == "14 instâncias"
    assert contagem("IP público", 1) == "1 IP público"
    assert contagem("Banco de dados", 3) == "3 bancos de dados"


def test_resposta_diz_o_bloco_de_cada_grupo_e_o_que_cada_unidade_tem() -> None:
    plano = _documento(incluir_antes=True, incluir_depois=True).plano
    blocos = {g.nome: g.bloco for g in plano.grupos_do_plano}
    assert blocos["catalogo"] == "pre"
    assert blocos["monitor"] == "pos"
    # Pré primeiro, pós por último: a ordem da tela
    assert (plano.grupos_do_plano[0].bloco, plano.grupos_do_plano[-1].bloco) == ("pre", "pos")
    papeis = {u.nome: (u.papel, u.motivo is not None) for u in plano.unidades_do_plano}
    assert papeis["checkout"] == ("escolhido", False)
    assert papeis["auth"] == ("vai_junto", True)
    assert papeis["catalogo"] == ("incluido", False)
    auth = next(u for u in plano.unidades_do_plano if u.nome == "auth")
    assert {(c.tipo, c.quantidade) for c in auth.conteudo} == {("instancia", 1), ("banco", 1)}
    # Os recursos vêm do mais caro para o mais barato
    assert [r.custo for r in auth.recursos] == sorted(
        (r.custo for r in auth.recursos), reverse=True
    )
    assert auth.recursos_total == len(auth.recursos)
    assert plano.etapas_da_conta >= 2
